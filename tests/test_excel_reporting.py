from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from excel_reporting import REPORT_GROUPS, _format_worklist, build_excel_download
from inventory_matching import MatchResult


def _result(system="장관계", kind="세균", validation="특이도"):
    return {
        "organism": "Escherichia coli",
        "relation": "근연종",
        "scope": "표적 직접 연관",
        "basis": "긴 검토 근거를 줄바꿈하여 읽기 쉽게 표시하는지 확인합니다.",
        "보유 여부": "보유",
        "보유 자원 수": 2,
        "관리번호": "Z001, Z002",
        "매칭 점수": 100.0,
        "사내 매칭명": "Escherichia coli strain A; Escherichia coli strain B",
        "Cat no.": "00123",
        "매칭 방식": "학명 일치",
        "원본 시트": "Sheet1",
        "system": system,
        "kind": kind,
        "input_targets": ("Target A", "Target B"),
        "검증 구분": validation,
        "우선순위": "필수",
    }


def test_excel_report_has_readable_alignment_borders_and_dimensions():
    results = [_result(system, kind) for _, system, kind in REPORT_GROUPS]
    results.append(_result(validation="포괄성"))
    report_bytes = build_excel_download(results)
    workbook = load_workbook(BytesIO(report_bytes))
    assert workbook.sheetnames == [
        "요약", "시험_작업목록", "후보_전체", "포괄성_후보",
        *(group[0] for group in REPORT_GROUPS),
    ]

    summary = workbook["요약"]
    assert summary["A1"].value == "qPCR CrossCheck · Validation Report"
    assert summary["A8"].value == len(results)
    assert summary["C8"].value == 1
    assert summary["E8"].value == len(REPORT_GROUPS)
    assert summary["G8"].value == len(results)
    assert "in-silico" in summary["B21"].value
    assert summary.sheet_view.showGridLines is False

    worklist = workbook["시험_작업목록"]
    assert worklist.auto_filter.ref == f"A1:V{len(results) + 1}"
    assert worklist["A1"].alignment.horizontal == "center"
    assert worklist["C2"].value == "장관계"
    assert worklist["D2"].value == "세균"
    assert worklist["S2"].value == "미보유"
    assert worklist["B2"].value == "Target A + Target B"
    assert worklist["E2"].value == "특이도"
    assert worklist["F2"].value == "필수"
    assert worklist["S2"].fill.fgColor.rgb.endswith("FEE2E2")
    assert worklist.freeze_panes == "G2"
    assert summary["B13"].value == 1
    assert summary["C13"].value == len(REPORT_GROUPS)
    assert summary["D13"].value == len(results)
    assert summary["G13"].value == 0
    assert summary["G15"].value == len(results)
    assert worklist.page_setup.orientation == "landscape"
    assert worklist.sheet_view.showGridLines is False
    assert len(worklist.tables) == 0

    for sheet in workbook.worksheets[2:]:
        assert sheet.freeze_panes == "D2"
        assert sheet.auto_filter.ref == sheet.dimensions
        assert not sheet.sheet_view.showGridLines
        assert sheet.row_dimensions[1].height == 34
        assert sheet.row_dimensions[2].height >= 30
        assert sheet["A1"].fill.fgColor.rgb.endswith("17324D")
        assert sheet["A1"].font.bold
        assert sheet["A1"].alignment.horizontal == "center"

        for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=14):
            for cell in row:
                assert cell.alignment.vertical == "center"
                assert cell.alignment.wrap_text
                assert cell.font.name == "맑은 고딕"
                assert cell.border.bottom.style in {"thin", "medium"}

        assert sheet["A2"].fill.fgColor.rgb.endswith("F8FAFC")
        assert sheet["I2"].number_format == "@"
        assert sheet["L2"].number_format == "@"
        assert sheet["H2"].number_format == "#,##0"
        assert sheet["J2"].number_format == "0.0"
        assert len(sheet.tables) == 0

    with ZipFile(BytesIO(report_bytes)) as archive:
        assert not any(name.startswith("xl/tables/") for name in archive.namelist())


def test_excel_summary_values_are_saved_for_viewers_without_formula_recalculation():
    results = [_result(), _result(validation="포괄성")]
    workbook = load_workbook(BytesIO(build_excel_download(results)), data_only=True)
    summary = workbook["요약"]

    assert summary["A8"].value == 2
    assert summary["C8"].value == 1
    assert summary["E8"].value == 1
    assert summary["G8"].value == 2
    assert [summary.cell(row, 2).value for row in range(13, 16)] == [1, 0, 0]
    assert [summary.cell(row, 3).value for row in range(13, 16)] == [1, 0, 0]
    assert [summary.cell(row, 4).value for row in range(13, 16)] == [2, 0, 0]
    assert [summary.cell(row, 7).value for row in range(13, 17)] == [0, 0, 2, 0]


def test_excel_readiness_summary_uses_same_required_recommended_scope_as_app_default():
    required = _result()
    reference = _result()
    reference["우선순위"] = "참고"

    workbook = load_workbook(BytesIO(build_excel_download([required, reference])), data_only=True)
    summary = workbook["요약"]

    assert [summary.cell(row, 7).value for row in range(13, 17)] == [0, 0, 1, 0]
    assert "필수·권장 작업 1건 기준" in summary["F17"].value
    assert "참고 1건 제외" in summary["F17"].value


def test_excel_report_adds_sheet_for_dynamically_inferred_pathogen_group():
    results = [_result("발열·매개체", "바이러스")]
    workbook = load_workbook(BytesIO(build_excel_download(results)))
    assert "발열·매개체_바이러스" in workbook.sheetnames
    assert workbook["발열·매개체_바이러스"]["A2"].value == "Escherichia coli"


def test_worklist_formatting_tolerates_an_optional_trailing_column_being_absent():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "qPCR 구성", "입력 표적", "검증 구분", "우선순위", "추천 미생물", "관계 분류",
        "선정 범위", "보유 여부", "관리번호", "사내 자원명", "Cat no.", "구매일",
        "최초 원액 용량 (µL)", "원액 누적 사용량 (µL)", "원액 잔량 (µL)",
        "희석액(1/100) 튜브 수 (n)", "준비 상태", "재고 점검", "비고",
    ])
    sheet.append(["Multiplex qPCR", "Listeria", "특이도", "필수", "Listeria innocua"] + [None] * 11 + ["미보유", "—", "—"])

    _format_worklist(sheet)

    assert sheet["Q2"].fill.fgColor.rgb.endswith("FEE2E2")


def test_excel_keeps_unprovided_stock_quantities_blank():
    result = _result()
    result["자원 상세"] = [MatchResult(
        True, 100.0, "정규화 일치", "Escherichia coli", "Z001", "00123", "Sheet1",
        remaining_volume_ul="소진", dilution_tubes="소진",
    )]

    workbook = load_workbook(BytesIO(build_excel_download([result])))
    sheet = workbook["시험_작업목록"]
    columns = {cell.value: cell.column for cell in sheet[1]}

    for header in (
        "최초 원액 용량 (µL)", "원액 누적 사용량 (µL)", "원액 잔량 (µL)",
        "희석액(1/100) 튜브 수 (n)",
    ):
        assert sheet.cell(2, columns[header]).value is None
    assert sheet.cell(2, columns["준비 상태"]).value == "소진"
