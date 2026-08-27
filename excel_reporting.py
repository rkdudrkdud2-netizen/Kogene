"""qPCR CrossCheck 결과 Excel 생성 및 전문 보고서 서식 지정."""

from __future__ import annotations

from collections import Counter
from io import BytesIO
import re

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from inventory_matching import excel_safe_value
from worklist import build_worklist_rows


REPORT_COLUMNS = {
    "organism": "미생물", "검증 구분": "검증 구분", "우선순위": "우선순위",
    "relation": "관계 분류", "scope": "선정 범위", "basis": "검토 근거",
    "보유 여부": "보유 여부", "보유 자원 수": "보유 자원 수", "관리번호": "관리번호",
    "매칭 점수": "최고 매칭 점수", "사내 매칭명": "사내 매칭명",
    "Cat no.": "Cat no.", "매칭 방식": "매칭 방식", "원본 시트": "원본 시트",
}

REPORT_GROUPS = (
    ("장관계_세균", "장관계", "세균"), ("장관계_바이러스", "장관계", "바이러스"),
    ("호흡기계_세균", "호흡기계", "세균"), ("호흡기계_바이러스", "호흡기계", "바이러스"),
    ("혈액매개_기생충", "혈액매개", "기생충"),
    ("기타_사용자입력", "기타", "미분류"),
)

FONT_NAME = "맑은 고딕"
NAVY, TEAL = "17324D", "0F766E"
TEAL_SOFT, AMBER_SOFT, SLATE_SOFT = "CCFBF1", "FEF3C7", "F1F5F9"
TEXT, MUTED, LINE, WHITE = "172033", "64748B", "DCE3EC", "FFFFFF"

COLUMN_WIDTHS = (30, 14, 12, 20, 18, 52, 12, 14, 30, 14, 44, 24, 16, 18)
WORKLIST_WIDTHS = (16, 28, 14, 12, 14, 12, 30, 20, 18, 12, 16, 40, 20, 16, 18, 20, 18, 22, 18, 30, 44, 14)
LEFT_HEADERS = {
    "입력 표적", "추천 미생물", "관계 분류", "선정 범위", "사내 자원명", "비고",
    "미생물", "검토 근거", "사내 매칭명", "매칭 방식", "원본 시트",
}
NUMERIC_HEADERS = {
    "보유 자원 수", "최고 매칭 점수", "매칭 점수", "최초 원액 용량 (µL)",
    "원액 누적 사용량 (µL)", "원액 잔량 (µL)", "희석액(1/100) 튜브 수 (n)",
}


def _estimated_row_height(row, widths) -> float:
    line_count = 1
    for cell, width in zip(row, widths):
        text = "" if cell.value is None else str(cell.value)
        line_count = max(line_count, max(1, (len(text) + max(1, int(width)) - 1) // max(1, int(width))))
    return min(78, max(27, line_count * 15))


def _add_priority_formatting(sheet, priority_column: int, last_row: int):
    if last_row < 2:
        return
    letter = get_column_letter(priority_column)
    target = f"{letter}2:{letter}{last_row}"
    for label, fill, color in (
        ("필수", TEAL_SOFT, "115E59"), ("권장", AMBER_SOFT, "92400E"), ("참고", SLATE_SOFT, "475569"),
    ):
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'${letter}2="{label}"'], fill=PatternFill("solid", fgColor=fill),
                font=Font(name=FONT_NAME, size=9, bold=True, color=color),
            ),
        )


def _style_sheet(sheet, widths, *, freeze="D2"):
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.tabColor = TEAL
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    sheet.oddHeader.center.text = f"&BqPCR CrossCheck  |  {sheet.title}"
    sheet.oddFooter.right.text = "Page &P / &N"
    sheet.oddFooter.left.text = "Generated report"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    headers = [cell.value for cell in sheet[1]]
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name=FONT_NAME, size=9, color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    sheet.row_dimensions[1].height = 34

    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        for column_index, cell in enumerate(row, start=1):
            header = headers[column_index - 1] if column_index <= len(headers) else ""
            horizontal = "left" if header in LEFT_HEADERS else "right" if header in NUMERIC_HEADERS else "center"
            cell.font = Font(name=FONT_NAME, size=9, color=TEXT)
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=LINE))
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
        sheet.row_dimensions[row_index].height = _estimated_row_height(row, widths)

    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_area = sheet.dimensions
    priority_column = next((index for index, value in enumerate(headers, start=1) if value == "우선순위"), None)
    if priority_column:
        _add_priority_formatting(sheet, priority_column, sheet.max_row)


def _format_candidate_numbers(sheet):
    for row in sheet.iter_rows(min_row=2):
        row[7].number_format = "#,##0"
        row[9].number_format = "0.0"
        row[8].number_format = "@"
        row[11].number_format = "@"


def _format_worklist(sheet):
    status_fills = {
        "즉시 사용 가능": "DCFCE7", "희석 필요": "FEF3C7", "원액 사용 가능": "FEF3C7",
        "소진": "FEE2E2", "미보유": "FEE2E2",
    }
    columns = {cell.value: cell.column - 1 for cell in sheet[1] if cell.value}
    text_headers = ("관리번호", "Cat no.")
    decimal_headers = ("최초 원액 용량 (µL)", "원액 누적 사용량 (µL)", "원액 잔량 (µL)")
    for row in sheet.iter_rows(min_row=2):
        for header in text_headers:
            if header in columns:
                row[columns[header]].number_format = "@"
        for header in decimal_headers:
            if header in columns:
                row[columns[header]].number_format = "#,##0.0"
        if "희석액(1/100) 튜브 수 (n)" in columns:
            row[columns["희석액(1/100) 튜브 수 (n)"]].number_format = "#,##0"
        if "매칭 점수" in columns:
            row[columns["매칭 점수"]].number_format = "0.0"
        status_index = columns.get("준비 상태")
        check_index = columns.get("재고 점검")
        status = str(row[status_index].value or "") if status_index is not None else ""
        if status in status_fills:
            row[status_index].fill = PatternFill("solid", fgColor=status_fills[status])
            row[status_index].font = Font(name=FONT_NAME, size=9, bold=True, color=TEXT)
        if check_index is not None and str(row[check_index].value or "") not in {"정상", "—"}:
            row[check_index].fill = PatternFill("solid", fgColor="FFEDD5")
            row[check_index].font = Font(name=FONT_NAME, size=9, bold=True, color="9A3412")


def _build_summary_counts(results, worklist_records):
    """수식 계산 엔진이 없는 뷰어에서도 보이도록 요약 숫자를 미리 계산한다."""
    validation_counts = Counter(str(item.get("검증 구분", "")) for item in results)
    priority_counts = Counter(
        (str(item.get("검증 구분", "")), str(item.get("우선순위", "")))
        for item in results
    )
    readiness_counts = Counter(str(row.get("준비 상태", "")) for row in worklist_records)
    return {
        "cards": {
            "전체 후보": len(results),
            "포괄성": validation_counts["포괄성"],
            "특이도": validation_counts["특이도"],
            "보유 후보": sum(str(item.get("보유 여부", "")) == "보유" for item in results),
        },
        "priorities": {
            priority: {
                "포괄성": priority_counts[("포괄성", priority)],
                "특이도": priority_counts[("특이도", priority)],
            }
            for priority in ("필수", "권장", "참고")
        },
        "readiness": {
            "즉시 사용 가능": readiness_counts["즉시 사용 가능"],
            "원액·희석 준비": readiness_counts["희석 필요"] + readiness_counts["원액 사용 가능"],
            "소진·미보유": readiness_counts["소진"] + readiness_counts["미보유"],
            "정보·재고 확인": readiness_counts["정보 확인 필요"],
        },
    }


def _add_summary_sheet(
    workbook, *, assay_type, target, summary_counts,
):
    sheet = workbook.create_sheet("요약", 0)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 95
    sheet.sheet_properties.tabColor = NAVY
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.35, right=0.35, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    for column, width in zip("ABCDEFGH", (18,) * 8):
        sheet.column_dimensions[column].width = width

    sheet["A1"] = "qPCR CrossCheck · Validation Report"
    sheet["A1"].font = Font(name=FONT_NAME, size=20, bold=True, color=WHITE)
    sheet["A1"].alignment = Alignment(horizontal="centerContinuous", vertical="center")
    for row in sheet["A1:H2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 18

    sheet["A3"] = "포괄성(Inclusivity)과 특이도(Exclusivity/Cross-reactivity) 후보 및 사내 자원 준비 현황"
    sheet["A3"].font = Font(name=FONT_NAME, size=10, color=MUTED)
    for cell in sheet["A3:H3"][0]:
        cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
    for row_index, (label, value) in enumerate((("qPCR 구성", assay_type or "qPCR"), ("입력 표적", target or "—")), start=4):
        sheet[f"A{row_index}"] = label
        sheet[f"A{row_index}"].font = Font(name=FONT_NAME, size=9, bold=True, color=TEAL)
        sheet[f"B{row_index}"] = value
        sheet[f"B{row_index}"].font = Font(name=FONT_NAME, size=9, color=TEXT)
        sheet[f"B{row_index}"].alignment = Alignment(horizontal="left", vertical="center")

    cards = (
        ("A7:B7", "A8:B9", "전체 후보", summary_counts["cards"]["전체 후보"], NAVY),
        ("C7:D7", "C8:D9", "포괄성", summary_counts["cards"]["포괄성"], TEAL),
        ("E7:F7", "E8:F9", "특이도", summary_counts["cards"]["특이도"], "155E75"),
        ("G7:H7", "G8:H9", "보유 후보", summary_counts["cards"]["보유 후보"], "0E7490"),
    )
    for label_range, value_range, label, value, color in cards:
        label_cell, value_cell = sheet[label_range.split(":")[0]], sheet[value_range.split(":")[0]]
        label_cell.value = label
        label_cell.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.value = value
        value_cell.number_format = "#,##0"
        value_cell.font = Font(name=FONT_NAME, size=19, bold=True, color=color)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet[label_range]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
        for row in sheet[value_range]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
                cell.border = Border(bottom=Side(style="thin", color=LINE))
                cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")

    sheet["A11"], sheet["F11"] = "우선순위별 후보", "작업 준비 상태"
    sheet["A11"].font = sheet["F11"].font = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
    for start, headers in ((1, ("우선순위", "포괄성", "특이도", "합계")), (6, ("상태", "작업 수", "표시"))):
        for column_index, value in enumerate(headers, start=start):
            cell = sheet.cell(12, column_index, value)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, (priority, fill, color) in enumerate(
        (("필수", TEAL_SOFT, "115E59"), ("권장", AMBER_SOFT, "92400E"), ("참고", SLATE_SOFT, "475569")), start=13,
    ):
        sheet.cell(row_index, 1, priority)
        inclusivity_count = summary_counts["priorities"][priority]["포괄성"]
        specificity_count = summary_counts["priorities"][priority]["특이도"]
        sheet.cell(row_index, 2, inclusivity_count)
        sheet.cell(row_index, 3, specificity_count)
        sheet.cell(row_index, 4, inclusivity_count + specificity_count)
        for column_index in range(2, 5):
            sheet.cell(row_index, column_index).number_format = "#,##0"
        for cell in sheet[row_index][0:4]:
            cell.font = Font(name=FONT_NAME, size=9, bold=cell.column == 1, color=color if cell.column == 1 else TEXT)
            cell.fill = PatternFill("solid", fgColor=fill if cell.column == 1 else WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="thin", color=LINE))

    readiness = (
        ("즉시 사용 가능", "DCFCE7", summary_counts["readiness"]["즉시 사용 가능"]),
        ("원액·희석 준비", AMBER_SOFT, summary_counts["readiness"]["원액·희석 준비"]),
        ("소진·미보유", "FEE2E2", summary_counts["readiness"]["소진·미보유"]),
        ("정보·재고 확인", "FFEDD5", summary_counts["readiness"]["정보·재고 확인"]),
    )
    for row_index, (status, fill, count) in enumerate(readiness, start=13):
        sheet.cell(row_index, 6, status)
        sheet.cell(row_index, 7, count)
        sheet.cell(row_index, 7).number_format = "#,##0"
        sheet.cell(row_index, 8, "●")
        for cell in sheet[row_index][5:8]:
            cell.font = Font(name=FONT_NAME, size=9, bold=cell.column == 8, color=TEXT)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        sheet.cell(row_index, 8).fill = PatternFill("solid", fgColor=fill)

    for cell in sheet["A18:H19"]:
        for item in cell:
            item.fill = PatternFill("solid", fgColor="EAF4F4")
    sheet["A18"], sheet["B18"], sheet["D18"], sheet["F18"] = (
        "읽는 순서", "① 포괄성 확인", "② 특이도 우선 검토", "③ 보유·준비 상태 확인",
    )
    for column in range(1, 9):
        cell = sheet.cell(18, column)
        cell.font = Font(name=FONT_NAME, size=9, bold=column == 1, color=NAVY if column == 1 else MUTED)
        cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
    sheet["A19"] = "표적 자체·동일 종/strain → 필수·권장 후보 → 시험 작업목록 순으로 검토"
    for cell in sheet["A19:H19"][0]:
        cell.font = Font(name=FONT_NAME, size=8, color=MUTED)
        cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")

    for row in sheet["A21:H22"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=AMBER_SOFT)
            cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
    sheet["A21"] = "주의"
    sheet["A21"].font = Font(name=FONT_NAME, size=9, bold=True, color="92400E")
    sheet["B21"] = "본 목록은 시험 설계 후보이며 실제 교차반응 여부는 in-silico 분석과 실험 결과로 확정해야 합니다."
    sheet["B21"].font = Font(name=FONT_NAME, size=9, italic=True, color="92400E")
    sheet.print_area = "A1:H22"
    sheet.oddFooter.right.text = "Page &P / &N"


def build_excel_download(results, assay_type: str = "", target: str = "") -> bytes:
    """요약 대시보드와 정돈된 포괄성·특이도·작업목록 시트를 생성한다."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        worklist_records = [
            {key: excel_safe_value(value) for key, value in row.items()}
            for row in build_worklist_rows(results, assay_type or "qPCR", target or "—")
        ]
        pd.DataFrame(worklist_records).to_excel(writer, sheet_name="시험_작업목록", index=False)
        worklist_sheet = writer.book["시험_작업목록"]
        _style_sheet(worklist_sheet, WORKLIST_WIDTHS, freeze="G2")
        _format_worklist(worklist_sheet)

        all_records = [
            {heading: excel_safe_value(item.get(key, "")) for key, heading in REPORT_COLUMNS.items()} for item in results
        ]
        all_frame = pd.DataFrame(all_records, columns=list(REPORT_COLUMNS.values()))
        all_frame.to_excel(writer, sheet_name="후보_전체", index=False)
        all_sheet = writer.book["후보_전체"]
        _style_sheet(all_sheet, COLUMN_WIDTHS, freeze="D2")
        _format_candidate_numbers(all_sheet)

        inclusivity_records = [
            {heading: excel_safe_value(item.get(key, "")) for key, heading in REPORT_COLUMNS.items()}
            for item in results if item.get("검증 구분") == "포괄성"
        ]
        pd.DataFrame(inclusivity_records, columns=list(REPORT_COLUMNS.values())).to_excel(writer, sheet_name="포괄성_후보", index=False)
        inclusivity_sheet = writer.book["포괄성_후보"]
        _style_sheet(inclusivity_sheet, COLUMN_WIDTHS, freeze="D2")
        _format_candidate_numbers(inclusivity_sheet)
        inclusivity_sheet.sheet_properties.tabColor = "14B8A6"

        report_groups = list(REPORT_GROUPS)
        known_groups = {(system, kind) for _, system, kind in report_groups}
        used_sheet_names = {name for name, _, _ in report_groups}
        for item in results:
            if item.get("검증 구분", "특이도") != "특이도":
                continue
            group = (item["system"], item["kind"])
            if group in known_groups:
                continue
            base_name = re.sub(r"[\\/*?:\[\]]", "_", f"{group[0]}_{group[1]}")[:31] or "추가_분류"
            sheet_name, suffix = base_name, 2
            while sheet_name in used_sheet_names:
                ending = f"_{suffix}"
                sheet_name = base_name[:31 - len(ending)] + ending
                suffix += 1
            report_groups.append((sheet_name, *group))
            known_groups.add(group)
            used_sheet_names.add(sheet_name)

        for sheet_name, system, kind in report_groups:
            records = [
                {heading: excel_safe_value(item.get(key, "")) for key, heading in REPORT_COLUMNS.items()}
                for item in results if item.get("검증 구분", "특이도") == "특이도"
                and item["system"] == system and item["kind"] == kind
            ]
            if not records:
                continue
            pd.DataFrame(records, columns=list(REPORT_COLUMNS.values())).to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.book[sheet_name]
            _style_sheet(sheet, COLUMN_WIDTHS, freeze="D2")
            _format_candidate_numbers(sheet)
            sheet.sheet_properties.tabColor = "155E75"

        _add_summary_sheet(
            writer.book, assay_type=assay_type, target=target,
            summary_counts=_build_summary_counts(results, worklist_records),
        )
        writer.book.calculation.calcMode = "auto"
        writer.book.calculation.fullCalcOnLoad = True
        writer.book.calculation.forceFullCalc = True
    return output.getvalue()
